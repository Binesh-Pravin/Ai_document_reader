"""
Read a bank statement PDF (downloaded from net banking, i.e. it contains real text).

Flow:
  1. Open the PDF (password if the bank protected it)
  2. Parse transaction lines with plain code (fast, no AI misreading digits)
  3. Work out debit / credit from how the running balance moves
  4. Verify: every row's balance must equal the previous balance +/- its amount
  5. Read account details (holder, account number, bank, branch) with the local LLM
  6. Export to Excel

Command line:
  python statement_extract.py statement.pdf
  python statement_extract.py statement.pdf --password 12345 --xlsx out.xlsx
"""
import io
import json
import re
import sys
from datetime import date
from decimal import Decimal, InvalidOperation

import pdfplumber
import requests
from pdfminer.pdfdocument import PDFPasswordIncorrect
from pdfplumber.utils.exceptions import PdfminerException

OLLAMA = "http://localhost:11434"
MODEL = "gemma3:12b"
TOL = Decimal("0.005")


class NeedPassword(Exception):
    """The PDF is protected and the password is missing or wrong."""


class NotTextPdf(Exception):
    """The PDF has no readable text (probably a scan)."""


# ---------- Small parsing helpers ----------
MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
MON = "|".join(MONTHS)
DATE_PAT = (rf"(?:\d{{1,2}}[/\-.]\d{{1,2}}[/\-.]\d{{2,4}}"
            rf"|\d{{1,2}}[ \-](?:{MON})[a-z]*[ \-,]+\d{{2,4}})")
LINE_RE = re.compile(rf"^\s*(?:\d{{1,4}}\s+)?({DATE_PAT})(?!\d)\s*(.*)$", re.I)
LEAD_DATE_RE = re.compile(rf"^({DATE_PAT})(?!\d)\s*", re.I)
TRAIL_DATE_RE = re.compile(rf"\s*{DATE_PAT}\s*$", re.I)
# an amount always has exactly two decimals; reference numbers do not, so they are ignored
AMOUNT_RE = re.compile(r"(?<![\d,])(?<!\d\.)(-?\d[\d,]*\.\d{2})(?![\d.])(?:\s*(CR|DR)\b)?", re.I)
OPENING_RE = re.compile(r"opening\s+balance|brought\s+forward|\bb/f\b", re.I)
SKIP_RE = re.compile(r"closing\s+balance|carried\s+forward|\bc/f\b|\btotal\b|statement\s+period", re.I)
NOISE_RE = re.compile(
    r"page\s*\d+\s*(of|/)\s*\d+|statement\s+of\s+account|opening\s+balance|closing\s+balance|"
    r"computer\s+generated|registered\s+office|customer\s+care|toll\s*free|end\s+of\s+statement|"
    r"\bdate\b.*\bbalance\b|narration.*balance|particulars.*balance|\bgstin\b|\bsummary\b", re.I)


def to_dec(text):
    text = text.replace(",", "").strip()
    neg = text.startswith("-") or (text.startswith("(") and text.endswith(")"))
    text = text.strip("()-")
    try:
        value = Decimal(text)
    except InvalidOperation:
        return None
    return -value if neg else value


def parse_date(text):
    s = re.sub(r"[ ,]+", " ", text.strip())
    m = re.match(r"^(\d{1,2})[/\-. ](\d{1,2})[/\-. ](\d{2,4})$", s)
    if m:
        d, mo, y = int(m[1]), int(m[2]), int(m[3])
    else:
        m = re.match(r"^(\d{1,2})[ \-]([A-Za-z]{3})[a-z]*[ \-](\d{2,4})$", s)
        if not m or m[2].lower() not in MONTHS:
            return None
        d, mo, y = int(m[1]), MONTHS[m[2].lower()], int(m[3])
    if y < 100:
        y += 2000
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def split_amounts(rest):
    """Return (description, candidate amounts, balance, direction hint) or None.
    The trailing group of adjacent amounts is [debit and/or credit] + balance."""
    ms = list(AMOUNT_RE.finditer(rest))
    if len(ms) < 2:
        return None
    group = [ms[-1]]
    for m in reversed(ms[:-1]):
        if len(group) == 3:
            break
        if rest[m.end():group[0].start()].strip() == "":
            group.insert(0, m)
        else:
            break
    if len(group) < 2:
        group = ms[-2:]

    desc = rest[:group[0].start()]
    bal_m = group[-1]
    balance = to_dec(bal_m.group(1))
    if balance is None:
        return None
    if (bal_m.group(2) or "").upper() == "DR":
        balance = -balance
    cands, hint = [], None
    for m in group[:-1]:
        v = to_dec(m.group(1))
        if v is not None and abs(v) > 0:
            cands.append(abs(v))
            if m.group(2):
                hint = "debit" if m.group(2).upper() == "DR" else "credit"
    return desc, cands, balance, hint


def parse_line(line):
    m = LINE_RE.match(line)
    if not m:
        return None
    d = parse_date(m.group(1))
    if d is None:
        return None
    rest = LEAD_DATE_RE.sub("", m.group(2), count=1)  # drop a second (value) date
    parts = split_amounts(rest)
    row = {"date": d, "description": rest.strip(), "cands": None, "balance": None, "hint": None}
    if parts:
        desc, cands, balance, hint = parts
        row.update(description=TRAIL_DATE_RE.sub("", desc).strip(), cands=cands, balance=balance, hint=hint)
    return row


HEADER_WORDS = {"date", "narration", "description", "particulars", "withdrawal", "withdrawals",
                "deposit", "deposits", "debit", "credit", "balance", "chq", "ref", "cheque",
                "amount", "amt", "value", "remarks", "txn", "transaction", "dt", "no"}


def is_noise(text):
    """Page footers, repeated column headings, totals: never part of a transaction."""
    if NOISE_RE.search(text) or re.search(r"\btotal\b", text, re.I):
        return True
    words = set(re.findall(r"[a-z]+", text.lower()))
    return len(words & HEADER_WORDS) >= 3


def page_blocks(page):
    """Group the page's text lines into blocks by vertical gap. Lines of one table row sit
    close together; rows are separated by a larger gap. This keeps wrapped descriptions with
    the right row, whether the date sits on the top, middle or bottom line of the row."""
    lines = page.extract_text_lines(return_chars=False) or []
    blocks, cur, prev = [], [], None
    for ln in lines:
        text = ln["text"].strip()
        if not text:
            continue
        if prev is not None and ln["top"] - prev["bottom"] > 0.5 * (prev["bottom"] - prev["top"]):
            blocks.append(cur)
            cur = []
        cur.append(text)
        prev = ln
    if cur:
        blocks.append(cur)
    return blocks


def row_from_block(block, idx, row):
    """One transaction spread over several lines of a block; idx is its date line."""
    pieces = []
    for i, text in enumerate(block):
        if i == idx:
            pieces.append(row["description"])
        elif is_noise(text) or (parse_line(text) and not parse_line(text)["description"]
                                and parse_line(text)["balance"] is None):
            continue   # footer / heading / a lone value-date line
        else:
            sp = split_amounts(text) if row["balance"] is None else None
            if sp:
                desc, row["cands"], row["balance"], row["hint"] = sp
                pieces.append(desc)
            else:
                pieces.append(text)
    row["description"] = " ".join(p for p in pieces if p).strip()
    return row


def parse_pages(pdf):
    rows, state = [], {"opening": None, "unparsed": 0}

    def accept(row):
        if row["balance"] is None:
            state["unparsed"] += 1
        elif OPENING_RE.search(row["description"]):
            state["opening"] = row["balance"]
        elif not SKIP_RE.search(row["description"]):
            rows.append(row)
            return row
        return None

    for pno, page in enumerate(pdf.pages, start=1):
        for block in page_blocks(page):
            found = [(i, parse_line(t)) for i, t in enumerate(block)]
            found = [(i, r) for i, r in found if r]
            strong = [(i, r) for i, r in found if r["balance"] is not None or r["description"]]
            starts = strong or found[:1]
            if len(starts) == 1:                      # the usual case: one block = one row
                i, row = starts[0]
                row["page"] = pno
                accept(row_from_block(block, i, row))
            elif len(starts) > 1:                     # rows packed with no gaps: top-aligned rows
                cur = pending = None
                for text in block:
                    row = parse_line(text)
                    if row:
                        row["page"] = pno
                        pending = cur = None
                        if row["balance"] is None:
                            pending = row
                        else:
                            cur = accept(row)
                    elif is_noise(text):
                        cur = pending = None
                    elif pending is not None:
                        sp = split_amounts(text)
                        if sp:
                            desc, pending["cands"], pending["balance"], pending["hint"] = sp
                            pending["description"] = (pending["description"] + " " + desc).strip()
                            cur, pending = accept(pending), None
                        else:
                            pending["description"] = (pending["description"] + " " + text).strip()
                    elif cur is not None:
                        cur["description"] += " " + text
    for r in rows:
        r["description"] = re.sub(r"\s+", " ", r["description"]).strip()
    return rows, state


# ---------- Verification ----------
def verify(rows, opening):
    """Decide debit/credit from balance movement and flag rows that do not reconcile."""
    reversed_order = len(rows) > 1 and rows[0]["date"] > rows[-1]["date"]
    if reversed_order:
        rows.reverse()  # newest-first statements: put them in time order

    prev = opening
    mismatches = []
    for i, r in enumerate(rows, start=1):
        r["row"] = i
        r["debit"] = r["credit"] = None
        bal, cands = r["balance"], r["cands"]
        r["status"], r["note"] = "ok", ""
        if prev is None:
            if cands and r["hint"]:
                r[r["hint"]] = cands[0]
            r["status"] = "unverified"
            r["note"] = "No opening balance found, so this first row cannot be checked."
            if cands and not r["hint"]:
                r["note"] += f" Amount {cands[0]:,.2f}: debit or credit unknown."
        else:
            delta = bal - prev
            match = next((c for c in cands if abs(abs(delta) - c) <= TOL), None)
            if match is None and not cands and abs(delta) <= TOL:
                match = Decimal(0)
            if match is not None:
                if delta >= 0:
                    r["credit"] = match if match else None
                else:
                    r["debit"] = match
            else:
                r["status"] = "mismatch"
                shown = f"{cands[0]:,.2f}" if cands else "none"
                r["note"] = (f"Balance moved by {abs(delta):,.2f} but the amount shown is {shown}. "
                             f"A digit may be misread or a transaction is missing before this row.")
                amount = cands[0] if cands else abs(delta)
                r["credit" if delta >= 0 else "debit"] = amount
                mismatches.append(r["row"])
        prev = bal

    debit_total = sum((r["debit"] or 0 for r in rows), Decimal(0))
    credit_total = sum((r["credit"] or 0 for r in rows), Decimal(0))
    closing = rows[-1]["balance"] if rows else None
    closing_ok = None
    if opening is not None and closing is not None:
        closing_ok = abs(opening + credit_total - debit_total - closing) <= TOL
    return {
        "rows": len(rows),
        "reversed_order": reversed_order,
        "mismatch_rows": mismatches,
        "unverified_rows": [r["row"] for r in rows if r["status"] == "unverified"],
        "debit_total": debit_total,
        "credit_total": credit_total,
        "opening_balance": opening,
        "closing_balance": closing,
        "closing_ok": closing_ok,
        "all_ok": bool(rows) and not mismatches and closing_ok is not False,
    }


# ---------- Account details (local LLM, every value checked against the page text) ----------
def read_header(text, model):
    warnings, header = [], {}
    flat = re.sub(r"\s+", " ", text)
    ifsc = re.search(r"\b[A-Z]{4}0[A-Z0-9]{6}\b", text)
    header["ifsc"] = ifsc.group(0) if ifsc else None

    prompt = (
        "Below is the first page of an Indian bank account statement. Return ONLY a JSON object "
        'with keys "account_holder", "account_number", "bank_name", "branch". '
        "Copy values exactly as printed. Use null if a value is not present. Do not guess.\n\n"
        + text[:3500])
    try:
        r = requests.post(f"{OLLAMA}/api/chat", timeout=300, json={
            "model": model, "stream": False, "format": "json",
            "messages": [{"role": "user", "content": prompt}],
            "options": {"temperature": 0, "num_ctx": 8192}})
        r.raise_for_status()
        data = json.loads(r.json()["message"]["content"])
    except (requests.RequestException, ValueError, KeyError):
        warnings.append("Account details were skipped because the local model could not be reached.")
        data = {}

    for key in ("account_holder", "account_number", "bank_name", "branch"):
        value = str(data.get(key) or "").strip() or None
        if value:  # keep only what really appears on the page
            if key == "account_number":
                ok = re.sub(r"\D", "", value) in re.sub(r"\D", "", flat)
            else:
                ok = re.sub(r"\s+", " ", value).lower() in flat.lower()
            if not ok:
                value = None
        header[key] = value
    return header, warnings


# ---------- Main entry ----------
def read_statement(path, password=None, model=MODEL):
    try:
        pdf = pdfplumber.open(path, password=password)
    except PDFPasswordIncorrect:
        raise NeedPassword("This PDF is password protected.")
    except PdfminerException as e:
        # pdfplumber wraps the original error; a missing or wrong password arrives this way
        inner = e.args[0] if e.args else None
        if isinstance(inner, PDFPasswordIncorrect):
            raise NeedPassword("This PDF is password protected.")
        raise NotTextPdf(f"This PDF could not be opened ({type(inner).__name__ if inner else 'unknown error'}).")
    with pdf:
        first_text = (pdf.pages[0].extract_text() or "") if pdf.pages else ""
        if len(first_text.strip()) < 40:
            raise NotTextPdf("No readable text found. This looks like a scanned statement.")
        rows, state = parse_pages(pdf)
        opening = state["opening"]
    if not rows:
        raise NotTextPdf("No transaction lines were recognised in this PDF.")

    header, warnings = read_header(first_text, model)
    if opening is None:
        m = OPENING_RE.search(first_text)
        if m:
            nums = AMOUNT_RE.search(first_text[m.end():m.end() + 80])
            opening = to_dec(nums.group(1)) if nums else None

    check = verify(rows, opening)
    if check["reversed_order"]:
        warnings.append("Transactions were listed newest first; they have been put in date order.")
    if state["unparsed"]:
        warnings.append(f"{state['unparsed']} line(s) started with a date but had no amounts and were skipped.")
    if check["unverified_rows"]:
        warnings.append("The opening balance was not found, so the first row could not be verified.")

    return {
        "header": header,
        "period": {"from": rows[0]["date"].isoformat(), "to": rows[-1]["date"].isoformat()},
        "verification": check,
        "transactions": rows,
        "warnings": warnings,
    }


def to_json_safe(result):
    """Decimals and dates to plain numbers/strings so the result can be sent as JSON."""
    def conv(v):
        if isinstance(v, Decimal):
            return float(round(v, 2))
        if isinstance(v, date):
            return v.isoformat()
        if isinstance(v, dict):
            return {k: conv(x) for k, x in v.items() if k not in ("cands", "hint")}
        if isinstance(v, (list, tuple)):
            return [conv(x) for x in v]
        return v
    return conv(result)


# ---------- Excel ----------
INDIAN = r"[>=10000000]##\,##\,##\,##0.00;[>=100000]##\,##\,##0.00;##,##0.00"


def to_excel(data):
    """data is the JSON-safe result. Returns the .xlsx file as bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    def style_header(ws):
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0F766E")

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    h, v, p = data["header"], data["verification"], data["period"]
    ws.append(["Item", "Value"])
    style_header(ws)
    for label, value in [
        ("Account holder", h.get("account_holder")), ("Account number", h.get("account_number")),
        ("Bank", h.get("bank_name")), ("Branch", h.get("branch")), ("IFSC", h.get("ifsc")),
        ("Period from", p["from"]), ("Period to", p["to"]),
        ("Opening balance", v.get("opening_balance")), ("Total debits", v["debit_total"]),
        ("Total credits", v["credit_total"]), ("Closing balance", v.get("closing_balance")),
        ("Transactions", v["rows"]),
        ("Verification", "All rows reconcile" if v["all_ok"]
         else f"{len(v['mismatch_rows'])} row(s) need checking"),
    ]:
        ws.append([label, "" if value is None else value])
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(horizontal="left")
        if isinstance(row[1].value, float):
            row[1].number_format = INDIAN
        elif isinstance(row[1].value, str):
            row[1].data_type = "s"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40

    tx = wb.create_sheet("Transactions")
    tx.append(["Date", "Description", "Debit", "Credit", "Balance", "Check"])
    style_header(tx)
    for r in data["transactions"]:
        d = date.fromisoformat(r["date"])
        status = {"ok": "OK", "mismatch": "CHECK", "unverified": "UNVERIFIED"}[r["status"]]
        tx.append([d, r["description"], r["debit"], r["credit"], r["balance"],
                   status + (": " + r["note"] if r["note"] else "")])
        row = tx[tx.max_row]
        row[0].number_format = "DD/MM/YYYY"
        for c in row[2:5]:
            c.number_format = INDIAN
        row[1].data_type = "s"   # never let text be read as a formula
        row[5].data_type = "s"
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        if r["status"] != "ok":
            for c in row:
                c.fill = PatternFill("solid", fgColor="FDE8D3")
    for col, width in zip("ABCDEF", (13, 60, 16, 16, 18, 50)):
        tx.column_dimensions[col].width = width
    tx.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(1)

    def opt(name):
        if name in args:
            i = args.index(name)
            val = args[i + 1]
            del args[i:i + 2]
            return val
        return None

    pw, xlsx, mdl = opt("--password"), opt("--xlsx"), opt("--model") or MODEL
    try:
        res = to_json_safe(read_statement(args[0], pw, mdl))
    except (NeedPassword, NotTextPdf) as e:
        sys.exit(str(e))
    v = res["verification"]
    print(json.dumps(res["header"], indent=2))
    print(f"\nPeriod {res['period']['from']} to {res['period']['to']}, {v['rows']} transactions")
    print("Verification:", "ALL ROWS RECONCILE" if v["all_ok"] else f"CHECK ROWS {v['mismatch_rows']}")
    for w in res["warnings"]:
        print("Note:", w)
    for r in res["transactions"][:5]:
        print(r["date"], r["debit"], r["credit"], r["balance"], "|", r["description"][:50])
    if xlsx:
        with open(xlsx, "wb") as f:
            f.write(to_excel(res))
        print("Saved", xlsx)
